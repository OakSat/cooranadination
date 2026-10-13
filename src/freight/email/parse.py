"""Raw MIME -> ParsedEmail (stdlib ``email``, no network, no LLM).

Two things matter downstream:
- ``new_text``: only what the sender wrote *this time*. Quoted history is cut off so the
  extractor never re-reads an old value as a fresh answer.
- ``dedupe_key``: the same email delivered twice (SES retry, S3 re-upload) gets the same key, so
  ingest processes it once.
"""

import csv
import hashlib
import io
import re
from datetime import UTC, datetime
from email import policy
from email.message import EmailMessage
from email.parser import BytesParser
from email.utils import getaddresses, parseaddr, parsedate_to_datetime
from html.parser import HTMLParser

from pydantic import BaseModel, ConfigDict, Field

MAX_ATTACHMENT_TEXT = 20_000  # characters per attachment fed to the extractor

# A line that starts quoted history, in the languages customers write in.
_REPLY_SEPARATORS = tuple(
    re.compile(p, re.IGNORECASE | re.MULTILINE)
    for p in (
        r"^On .{0,200}wrote:\s*$",
        r"^Am .{0,200}schrieb .{0,200}:\s*$",
        r"^El .{0,200}escribi[óo]:\s*$",
        r"^-{2,}\s*(Original Message|Ursprüngliche Nachricht|Mensaje original)\s*-{2,}",
        r"^(From|Von|De):\s.+\n(Sent|Gesendet|Enviado|Date|Datum|Fecha):\s",
        r"^_{10,}\s*$",  # Outlook's rule line
    )
)


class Attachment(BaseModel):
    model_config = ConfigDict(frozen=True)

    filename: str
    content_type: str
    size: int
    text: str = ""  # extracted text, empty for unsupported types
    data: bytes = Field(repr=False, exclude=True)


class ParsedEmail(BaseModel):
    model_config = ConfigDict(frozen=True)

    message_id: str
    in_reply_to: str | None
    references: list[str]
    from_addr: str
    from_name: str
    to: list[str]
    subject: str
    sent_at: datetime
    text: str  # full plain-text body
    new_text: str  # body without quoted history
    attachments: list[Attachment]
    dedupe_key: str

    @property
    def looks_like_reply(self) -> bool:
        has_prefix = re.match(r"^\s*(re|aw|antw|rv|r|wg|fw|fwd)\s*:", self.subject, re.I)
        return bool(self.in_reply_to or self.references or has_prefix or self.new_text != self.text)


class _TextExtractor(HTMLParser):
    _BLOCK = {"p", "div", "br", "tr", "li", "h1", "h2", "h3", "h4", "table", "blockquote"}

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in ("script", "style", "head"):
            self._skip += 1
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in ("script", "style", "head"):
            self._skip = max(0, self._skip - 1)
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skip:
            self.parts.append(data)


def html_to_text(html: str) -> str:
    p = _TextExtractor()
    p.feed(html)
    text = "".join(p.parts)
    return re.sub(r"\n{3,}", "\n\n", re.sub(r"[ \t]+", " ", text)).strip()


def split_reply(text: str) -> str:
    """Return the part of ``text`` above the first quoted-history marker."""
    cut = len(text)
    for pattern in _REPLY_SEPARATORS:
        if m := pattern.search(text):
            cut = min(cut, m.start())
    kept = [line for line in text[:cut].splitlines() if not line.lstrip().startswith(">")]
    return "\n".join(kept).strip()


def attachment_text(filename: str, content_type: str, data: bytes) -> str:
    name = filename.lower()
    try:
        if content_type == "application/pdf" or name.endswith(".pdf"):
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif name.endswith(".xlsx"):
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            rows = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        rows.append(" | ".join(cells))
            text = "\n".join(rows)
        elif name.endswith(".csv") or content_type == "text/csv":
            decoded = data.decode("utf-8-sig", errors="replace")
            text = "\n".join(" | ".join(r) for r in csv.reader(io.StringIO(decoded)) if r)
        elif content_type.startswith("text/"):
            decoded = data.decode("utf-8", errors="replace")
            text = html_to_text(decoded) if content_type == "text/html" else decoded
        else:
            return ""
    except Exception:  # a corrupt attachment must not sink the email
        return ""
    return text[:MAX_ATTACHMENT_TEXT]


def _body(msg: EmailMessage) -> str:
    plain = msg.get_body(preferencelist=("plain",))
    if plain is not None:
        return str(plain.get_content()).strip()
    html = msg.get_body(preferencelist=("html",))
    if html is not None:
        return html_to_text(str(html.get_content()))
    return ""


def _ids(value: str | None) -> list[str]:
    return re.findall(r"<[^<>\s]+>", value or "")


def _normalize_body(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip().casefold()


def parse_raw(raw: bytes) -> ParsedEmail:
    msg = BytesParser(policy=policy.default).parsebytes(raw)
    assert isinstance(msg, EmailMessage)
    from_name, from_addr = parseaddr(str(msg.get("From", "")))
    date_header = msg.get("Date")
    try:
        sent_at = parsedate_to_datetime(str(date_header)) if date_header else datetime.now(UTC)
    except (TypeError, ValueError):
        sent_at = datetime.now(UTC)
    if sent_at.tzinfo is None:
        sent_at = sent_at.replace(tzinfo=UTC)

    text = _body(msg)
    attachments = []
    for part in msg.iter_attachments():
        assert isinstance(part, EmailMessage)
        data = part.get_payload(decode=True)
        if not isinstance(data, bytes):
            continue
        filename = part.get_filename() or "attachment"
        ctype = part.get_content_type()
        attachments.append(
            Attachment(
                filename=filename,
                content_type=ctype,
                size=len(data),
                text=attachment_text(filename, ctype, data),
                data=data,
            )
        )

    message_id = (_ids(str(msg.get("Message-ID", ""))) or [""])[0]
    # Without a Message-ID, the same sender + timestamp + body is the same email.
    basis = message_id or f"{from_addr.lower()}|{sent_at.isoformat()}|{_normalize_body(text)}"
    in_reply = _ids(str(msg.get("In-Reply-To", "")))
    return ParsedEmail(
        message_id=message_id,
        in_reply_to=in_reply[0] if in_reply else None,
        references=_ids(str(msg.get("References", ""))),
        from_addr=from_addr.lower(),
        from_name=from_name,
        to=[a.lower() for _, a in getaddresses([str(v) for v in msg.get_all("To", [])])],
        subject=str(msg.get("Subject", "")).strip(),
        sent_at=sent_at,
        text=text,
        new_text=split_reply(text),
        attachments=attachments,
        dedupe_key=hashlib.sha256(basis.encode()).hexdigest(),
    )
